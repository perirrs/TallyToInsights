from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Body
from sqlalchemy.orm import Session
from app.database import get_db
from app.utils.auth import get_current_user
from app.models.user import User
import json
import os

router = APIRouter(prefix="/api/audit-checks", tags=["audit-checks"])

CHECKS_FILE = os.path.join(os.path.dirname(__file__), "../../data/checks_735.json")
OVERRIDES_FILE = os.path.join(os.path.dirname(__file__), "../../data/check_overrides.json")


def load_checks():
    if os.path.exists(CHECKS_FILE):
        with open(CHECKS_FILE) as f:
            return json.load(f)
    return []


def load_overrides():
    if os.path.exists(OVERRIDES_FILE):
        with open(OVERRIDES_FILE) as f:
            return json.load(f)
    return {"created": [], "updated": {}, "deleted": []}


def save_overrides(overrides):
    os.makedirs(os.path.dirname(OVERRIDES_FILE), exist_ok=True)
    with open(OVERRIDES_FILE, "w") as f:
        json.dump(overrides, f, indent=2)


def merge_checks():
    base = load_checks()
    overrides = load_overrides()
    deleted_ids = set(overrides.get("deleted", []))
    updated = overrides.get("updated", {})
    created = overrides.get("created", [])

    result = []
    for check in base:
        if check["id"] in deleted_ids:
            continue
        key = str(check["id"])
        if key in updated:
            result.append({**check, **updated[key]})
        else:
            result.append(check)

    result.extend(created)
    return result


@router.get("/")
def list_checks(
    search: str = "",
    category: str = "",
    risk: str = "",
    page: int = 1,
    page_size: int = 50,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    checks = merge_checks()

    if search:
        s = search.lower()
        checks = [c for c in checks if s in c.get("desc", "").lower()]
    if category:
        checks = [c for c in checks if c.get("category", "") == category]
    if risk:
        checks = [c for c in checks if c.get("risk", "") == risk]

    total = len(checks)
    start = (page - 1) * page_size
    items = checks[start: start + page_size]

    categories = sorted(set(c.get("category", "") for c in merge_checks() if c.get("category")))

    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "categories": categories,
    }


@router.post("/", status_code=201)
def create_check(data: dict, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    overrides = load_overrides()
    # Generate a new ID
    all_checks = merge_checks()
    max_id = max((c["id"] for c in all_checks), default=0)
    new_check = {
        "id": max_id + 1,
        "desc": data.get("desc", ""),
        "category": data.get("category", ""),
        "risk": data.get("risk", "Medium"),
        "feasibility": data.get("feasibility", "Auto"),
        "analysis": data.get("analysis", ""),
        "module": data.get("module", ""),
        "source": data.get("source", ""),
        "active": data.get("active", True),
    }
    overrides.setdefault("created", []).append(new_check)
    save_overrides(overrides)
    return new_check


@router.put("/{check_id}")
def update_check(check_id: int, data: dict, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    overrides = load_overrides()
    # Check if it's a created check
    created = overrides.get("created", [])
    for i, c in enumerate(created):
        if c["id"] == check_id:
            overrides["created"][i] = {**c, **data, "id": check_id}
            save_overrides(overrides)
            return overrides["created"][i]

    # Otherwise update base check
    base = load_checks()
    base_check = next((c for c in base if c["id"] == check_id), None)
    if not base_check:
        raise HTTPException(404, "Check not found")

    overrides.setdefault("updated", {})[str(check_id)] = data
    save_overrides(overrides)
    return {**base_check, **data}


@router.delete("/{check_id}", status_code=204)
def delete_check(check_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    overrides = load_overrides()
    # Remove from created if present
    created = overrides.get("created", [])
    new_created = [c for c in created if c["id"] != check_id]
    if len(new_created) < len(created):
        overrides["created"] = new_created
        save_overrides(overrides)
        return

    # Mark as deleted for base checks
    deleted = overrides.get("deleted", [])
    if check_id not in deleted:
        deleted.append(check_id)
        overrides["deleted"] = deleted
        save_overrides(overrides)


@router.post("/bulk-delete", status_code=200)
def bulk_delete_checks(
    data: dict = Body(...),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Delete multiple checks by ID list."""
    ids: list[int] = data.get("ids", [])
    if not ids:
        return {"deleted": 0}

    overrides = load_overrides()
    created = overrides.get("created", [])
    created_ids = {c["id"] for c in created}
    deleted_set = set(overrides.get("deleted", []))

    new_created = [c for c in created if c["id"] not in ids]
    for cid in ids:
        if cid not in created_ids and cid not in deleted_set:
            deleted_set.add(cid)

    overrides["created"] = new_created
    overrides["deleted"] = list(deleted_set)
    save_overrides(overrides)
    return {"deleted": len(ids)}


RISK_MAP = {"h": "High", "m": "Medium", "l": "Low", "high": "High", "medium": "Medium", "low": "Low"}
FEASIBILITY_MAP = {"auto": "Auto", "upload": "Upload", "module": "Module", "manual": "Manual"}


def _find_header_row(ws):
    """Find the row index (1-based) that contains the actual column headers.
    Looks for a row where a cell contains 'audit check description' or 'description'.
    Returns (header_row_index, headers_dict mapping normalized_name→col_index).
    """
    for row_idx in range(1, min(10, ws.max_row + 1)):
        row_vals = [str(cell.value or "").strip() for cell in ws[row_idx]]
        normalized = [v.lower() for v in row_vals]
        if any("audit check description" in v or (v in ("description", "desc")) for v in normalized):
            return row_idx, {v: i for i, v in enumerate(normalized) if v}
    return None, {}


def _col(row_vals, headers, *candidates):
    """Return the first matching cell value from a list of candidate header names."""
    for c in candidates:
        if c in headers:
            v = row_vals[headers[c]]
            if v is not None:
                return str(v).strip()
    return ""


@router.post("/upload-excel")
async def upload_excel(
    file: UploadFile = File(...),
    replace: bool = False,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    try:
        import openpyxl
        import io

        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active

        header_row_idx, headers = _find_header_row(ws)
        if header_row_idx is None:
            raise HTTPException(400, "Could not find header row. Expected a column named 'Audit Check Description'.")

        overrides = load_overrides()

        if replace:
            # Replace all: clear created/deleted overrides, rewrite the base file
            new_base = []
            serial = 1
            for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
                row_vals = [str(v).strip() if v is not None else "" for v in row]
                if not any(row_vals):
                    continue
                desc = _col(row_vals, headers, "audit check description", "description", "desc")
                if not desc:
                    continue
                risk_raw = _col(row_vals, headers, "risk").upper()
                risk = RISK_MAP.get(risk_raw.lower(), "Medium")
                feasibility_raw = _col(row_vals, headers, "feasibility").lower()
                feasibility = FEASIBILITY_MAP.get(feasibility_raw, "Auto")
                analysis = _col(row_vals, headers, "analysis type", "analysis")
                new_base.append({
                    "id": serial,
                    "desc": desc,
                    "category": _col(row_vals, headers, "category"),
                    "risk": risk,
                    "feasibility": feasibility,
                    "analysis": analysis,
                    "module": _col(row_vals, headers, "tally module", "module"),
                    "source": _col(row_vals, headers, "tally data source", "source"),
                    "active": True,
                })
                serial += 1
            os.makedirs(os.path.dirname(CHECKS_FILE), exist_ok=True)
            with open(CHECKS_FILE, "w") as f:
                json.dump(new_base, f, indent=2)
            # Clear overrides since base is fresh
            save_overrides({"created": [], "updated": {}, "deleted": []})
            return {"message": f"Replaced all checks — {len(new_base)} checks imported"}

        # Add mode: append to created overrides
        all_checks = merge_checks()
        max_id = max((c["id"] for c in all_checks), default=0)
        added = 0

        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            row_vals = [str(v).strip() if v is not None else "" for v in row]
            if not any(row_vals):
                continue
            desc = _col(row_vals, headers, "audit check description", "description", "desc")
            if not desc:
                continue
            risk_raw = _col(row_vals, headers, "risk")
            risk = RISK_MAP.get(risk_raw.lower(), "Medium")
            feasibility_raw = _col(row_vals, headers, "feasibility").lower()
            feasibility = FEASIBILITY_MAP.get(feasibility_raw, "Auto")
            analysis = _col(row_vals, headers, "analysis type", "analysis")
            max_id += 1
            overrides.setdefault("created", []).append({
                "id": max_id,
                "desc": desc,
                "category": _col(row_vals, headers, "category"),
                "risk": risk,
                "feasibility": feasibility,
                "analysis": analysis,
                "module": _col(row_vals, headers, "tally module", "module"),
                "source": _col(row_vals, headers, "tally data source", "source"),
                "active": True,
            })
            added += 1

        save_overrides(overrides)
        return {"message": f"Imported {added} checks successfully"}
    except HTTPException:
        raise
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")
    except Exception as e:
        raise HTTPException(400, f"Failed to parse Excel file: {str(e)}")
